import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import cv2
import fitz
import numpy as np
import openpyxl

from src.models import Box, Field, TemplatePreset
from src.processor import (
    generate_ui_templates,
    generate_ui_templates_multi,
    remap_preset_to_detected_layout,
    run_analysis,
    _validate_analysis_page_geometry,
)
from src.vision import (
    auto_detect_checkboxes,
    estimate_deskew_angle,
    load_pdf_pages,
)


_RUN_CORPUS_TESTS = os.getenv("RUN_SCAN_CORPUS_TESTS") == "1"
_RUN_TRIO_CORPUS_TESTS = (
    os.getenv("RUN_TRIO_CORPUS_TESTS") == "1" or _RUN_CORPUS_TESTS
)
_SCAN_DIR = Path(
    os.getenv("SURVEY_SCAN_CORPUS", r"C:\Users\Public\scan")
)
_PRESET_DIR = Path(
    os.getenv(
        "SURVEY_PRESET_DIR",
        str(Path(os.getenv("LOCALAPPDATA", "")) / "CheckFinder" / "presets"),
    )
)

_ONE_PAGE_PDFS = (
    "보드게임.pdf",
    "색깔.pdf",
    "운영실무.pdf",
    "웰다잉.pdf",
    "자활창업.pdf",
    "리더.pdf",
    "마음.pdf",
    "여행.pdf",
)
_TWO_PAGE_PDFS = (
    "거점.pdf",
    "급식.pdf",
    "도서관.pdf",
    "라라.pdf",
    "매점.pdf",
    "물류.pdf",
    "세차분식.pdf",
    "청소.pdf",
    "청수.pdf",
    "카페주거.pdf",
    "편의점.pdf",
    "헤이클린.pdf",
)

_ADDITIONAL_CORPUS_PDFS = (
    "S25C-0i26090316340.pdf",
    "법정의무.pdf",
    "법정의무2.pdf",
)

_INDEPENDENT_LAYOUT_TRIO = ("리더.pdf", "마음.pdf", "여행.pdf")

_TWO_PAGE_ROW_SIGNATURES = (
    (2, 5, 4, 4, 2, 2, 4, 2, 2, 3, 3, 2, 2),
    (2, 1, 4, 1, 1, 1, 1, 1, 1, 2, 2, 1, 4, 4),
)


def _group_boxes_by_row(boxes):
    rows = []
    current = []
    for box in sorted(boxes, key=lambda item: item[1]):
        if not current:
            current.append(box)
            continue
        previous = current[-1]
        tolerance = max(box[3], previous[3]) * 0.5
        if abs(box[1] - previous[1]) <= tolerance:
            current.append(box)
        else:
            rows.append(current)
            current = [box]
    if current:
        rows.append(current)
    return rows


def _representative_groups(paths):
    groups = []
    for size in (2, 3):
        for start in range(0, len(paths), size):
            group = tuple(paths[start : start + size])
            if len(group) < 2:
                group = tuple(paths[-size:])
            if group not in groups:
                groups.append(group)
    all_paths = tuple(paths)
    if all_paths not in groups:
        groups.append(all_paths)
    return groups


def _detected_layout_config(templates, page_count: int) -> TemplatePreset:
    """Build a disposable preset so remapping checks every local frame."""
    fields = []
    for page_idx in range(page_count):
        boxes = [
            Box(page_idx, x, y, w, h)
            for x, y, w, h in auto_detect_checkboxes(templates[page_idx])
        ]
        fields.append(Field(f"page_{page_idx + 1}", boxes=boxes))
    return TemplatePreset(page_count=page_count, fields=fields)


def _assert_review_grid_alignment(test_case, document, page_idx: int) -> None:
    """Check vector answer-cell overlays against frames in the embedded page."""
    page = document[page_idx]
    images = page.get_images(full=True)
    test_case.assertTrue(images, "review page has no embedded scan image")
    encoded = document.extract_image(images[0][0])["image"]
    gray = cv2.imdecode(np.frombuffer(encoded, np.uint8), cv2.IMREAD_GRAYSCALE)
    test_case.assertIsNotNone(gray)
    detected = auto_detect_checkboxes(cv2.cvtColor(gray, cv2.COLOR_GRAY2BGR))
    cells = [box for box in detected if 120 < box[2] < 240 and 90 < box[3] < 210]
    overlays = [
        drawing["rect"]
        for drawing in page.get_drawings()
        if 120 < drawing["rect"].width < 240
        and 90 < drawing["rect"].height < 210
    ]
    # Some forms also contain large non-answer frames.  The 25 configured
    # answer overlays must each select a distinct nearby detected frame.
    test_case.assertGreaterEqual(len(cells), 25)
    test_case.assertEqual(len(overlays), 25)
    edge_errors = []
    matched_cells = set()
    for rect in overlays:
        matched_idx, (x, y, w, h) = min(
            enumerate(cells),
            key=lambda indexed: (
                indexed[1][0] + indexed[1][2] / 2 - (rect.x0 + rect.x1) / 2
            ) ** 2
            + (
                indexed[1][1] + indexed[1][3] / 2 - (rect.y0 + rect.y1) / 2
            ) ** 2,
        )
        matched_cells.add(matched_idx)
        edge_errors.extend((x - rect.x0, y - rect.y0, x + w - rect.x1, y + h - rect.y1))
    test_case.assertEqual(len(matched_cells), 25)
    test_case.assertLessEqual(float(np.percentile(np.abs(edge_errors), 95)), 5.0)


@unittest.skipUnless(
    _RUN_CORPUS_TESTS,
    "set RUN_SCAN_CORPUS_TESTS=1 to run the external PDF corpus",
)
class ScanCorpusDetectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not _SCAN_DIR.is_dir():
            raise AssertionError(f"scan corpus directory is missing: {_SCAN_DIR}")
        expected = (
            set(_ONE_PAGE_PDFS)
            | set(_TWO_PAGE_PDFS)
            | set(_ADDITIONAL_CORPUS_PDFS)
        )
        actual = {path.name for path in _SCAN_DIR.glob("*.pdf")}
        if actual != expected:
            missing = sorted(expected - actual)
            unclassified = sorted(actual - expected)
            raise AssertionError(
                f"scan corpus changed; missing={missing}, unclassified={unclassified}"
            )

    @staticmethod
    def _page_angles(path: Path, page_count: int):
        pages = load_pdf_pages(
            str(path),
            page_indices=list(range(page_count)),
        )
        return [float(estimate_deskew_angle(page) or 0.0) for page in pages]

    def _assert_layout(self, templates, page_count: int):
        self.assertEqual(set(templates), set(range(page_count)))
        detected_by_page = {
            page: auto_detect_checkboxes(template)
            for page, template in templates.items()
        }

        if page_count == 1:
            rows = _group_boxes_by_row(detected_by_page[0])
            small_rows = [
                row
                for row in rows
                if np.median([box[2] for box in row]) <= 50
                and np.median([box[3] for box in row]) <= 50
            ]
            grid_rows = [
                row
                for row in rows
                if len(row) >= 3
                and np.median([box[2] for box in row]) >= 100
                and np.median([box[3] for box in row]) >= 100
            ]
            self.assertEqual([len(row) for row in small_rows], [2, 6, 4])
            self.assertEqual([len(row) for row in grid_rows], [5] * 6)
            return

        signatures = tuple(
            tuple(len(row) for row in _group_boxes_by_row(detected_by_page[page]))
            for page in range(page_count)
        )
        self.assertEqual(signatures, _TWO_PAGE_ROW_SIGNATURES)

    def test_every_pdf_individually(self):
        for page_count, names in (
            (1, _ONE_PAGE_PDFS),
            (2, _TWO_PAGE_PDFS),
        ):
            for name in names:
                with self.subTest(pdf=name, page_count=page_count):
                    path = _SCAN_DIR / name
                    angles = self._page_angles(path, page_count)
                    templates = generate_ui_templates(
                        str(path),
                        page_count,
                        -1,
                        0.0,
                        page_fine_angles=angles,
                    )
                    _validate_analysis_page_geometry(
                        [str(path)],
                        [templates[page] for page in range(page_count)],
                        TemplatePreset(page_count=page_count),
                    )
                    self._assert_layout(templates, page_count)

    def test_compatible_pdfs_in_pairs_triples_and_all(self):
        for page_count, names in (
            (1, _ONE_PAGE_PDFS),
            (2, _TWO_PAGE_PDFS),
        ):
            paths = [_SCAN_DIR / name for name in names]
            for group in _representative_groups(paths):
                labels = tuple(path.name for path in group)
                with self.subTest(pdfs=labels, page_count=page_count):
                    # Multi-file editing intentionally keeps the first file's
                    # pixels as the UI basis.  Analysis must instead build a
                    # reference for every file and remap the chosen fields to
                    # that local layout; checking every file against the first
                    # template would hide a coordinate regression.
                    source_templates = generate_ui_templates_multi(
                        [str(path) for path in group],
                        page_count,
                        -1,
                        0.0,
                        page_fine_angles=self._page_angles(group[0], page_count),
                    )
                    source_config = _detected_layout_config(
                        source_templates, page_count
                    )
                    for path in group:
                        local_templates = generate_ui_templates(
                            str(path),
                            page_count,
                            -1,
                            0.0,
                            page_fine_angles=self._page_angles(path, page_count),
                        )
                        _validate_analysis_page_geometry(
                            [str(path)],
                            [local_templates[page] for page in range(page_count)],
                            TemplatePreset(page_count=page_count),
                        )
                        remap = remap_preset_to_detected_layout(
                            source_config,
                            local_templates,
                            source_templates=source_templates,
                        )
                        self.assertTrue(remap.compatible)
                        self.assertTrue(remap.accepted)
                        self._assert_layout(local_templates, page_count)

    def test_one_page_pdfs_accept_default_preset_with_current_box_geometry(self):
        preset_path = _PRESET_DIR / "기본.json"
        template_path = _PRESET_DIR / "기본_tpl_p0.png"
        if not preset_path.is_file() or not template_path.is_file():
            raise AssertionError(
                f"default preset fixture is missing: {_PRESET_DIR}"
            )

        data = json.loads(preset_path.read_text(encoding="utf-8"))
        config = TemplatePreset(
            page_count=int(data.get("page_count", 1)),
            fine_angle=float(data.get("fine_angle", 0.0)),
            rot_code=int(data.get("rot_code", -1)),
            reverse_numbering=bool(data.get("reverse_numbering", True)),
            template_dilate_pct=float(data.get("template_dilate_pct", 0.3)),
            fields=[Field.from_dict(field) for field in data.get("fields", [])],
            page_fine_angles=[
                float(value) for value in data.get("page_fine_angles", [])
            ],
        )
        source_template = cv2.imdecode(
            np.frombuffer(template_path.read_bytes(), np.uint8),
            cv2.IMREAD_COLOR,
        )
        self.assertIsNotNone(source_template)

        color_pdf_path = _SCAN_DIR / "색깔.pdf"
        color_angles = []
        color_detected = []
        color_mapped_geometry = set()
        for name in _ONE_PAGE_PDFS:
            with self.subTest(pdf=name):
                pdf_path = _SCAN_DIR / name
                angles = self._page_angles(pdf_path, 1)
                templates = generate_ui_templates(
                    str(pdf_path),
                    1,
                    config.rot_code,
                    config.fine_angle,
                    page_fine_angles=angles,
                )
                detected = [
                    Box(0, x, y, w, h)
                    for x, y, w, h in auto_detect_checkboxes(templates[0])
                ]

                result = remap_preset_to_detected_layout(
                    config,
                    templates,
                    source_templates={0: source_template},
                    detected_boxes_by_page={0: detected},
                )

                self.assertTrue(result.accepted)
                self.assertTrue(result.compatible)
                self.assertEqual(result.matched_boxes, result.expected_boxes)
                detected_geometry = {
                    (box.page_idx, box.x, box.y, box.w, box.h)
                    for box in detected
                }
                mapped_geometry = {
                    (box.page_idx, box.x, box.y, box.w, box.h)
                    for field in result.config.fields
                    if not field.is_comment
                    for box in field.boxes
                }
                self.assertTrue(mapped_geometry <= detected_geometry)
                if pdf_path == color_pdf_path:
                    color_angles = angles
                    color_detected = detected
                    color_mapped_geometry = mapped_geometry

        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PyQt6.QtWidgets import QApplication

        from src.ui import MainWindow

        app = QApplication.instance() or QApplication([])
        window = MainWindow()
        try:
            window.preset_dir = _PRESET_DIR
            window.file_paths = [str(color_pdf_path)]
            window.pages = load_pdf_pages(
                str(color_pdf_path), page_indices=[0]
            )
            window.preset = TemplatePreset(page_count=1)
            window._pages_are_canonical = False
            window._reset_state_for_new_pdf()
            window.preset.page_fine_angles = color_angles
            window._update_page_size()
            raw_page = window.pages[0]
            window.auto_detect(mark_dirty=False)
            self.assertEqual(
                sum(len(field.boxes) for field in window.preset.fields),
                len(color_detected),
            )
            self.assertIs(window.pages[0], raw_page)
            self.assertIs(
                window._canvas_base_page(raw_page, 0),
                window._inferred_display_templates[0],
            )

            window._apply_loaded_preset(data, preset_name="기본")

            self.assertEqual(
                [field.name for field in window.preset.fields],
                [field.name for field in config.fields],
            )
            gui_geometry = {
                (box.page_idx, box.x, box.y, box.w, box.h)
                for field in window.preset.fields
                if not field.is_comment
                for box in field.boxes
            }
            self.assertEqual(gui_geometry, color_mapped_geometry)
        finally:
            window.close()
            app.processEvents()

    def test_single_response_keeps_answers_with_clean_preset_fallback(self):
        preset_path = _PRESET_DIR / "기본.json"
        template_path = _PRESET_DIR / "기본_tpl_p0.png"
        data = json.loads(preset_path.read_text(encoding="utf-8"))
        config = TemplatePreset(
            page_count=int(data.get("page_count", 1)),
            fine_angle=float(data.get("fine_angle", 0.0)),
            rot_code=int(data.get("rot_code", -1)),
            reverse_numbering=bool(data.get("reverse_numbering", True)),
            template_dilate_pct=float(data.get("template_dilate_pct", 0.3)),
            fields=[Field.from_dict(field) for field in data.get("fields", [])],
            page_fine_angles=[
                float(value) for value in data.get("page_fine_angles", [])
            ],
        )
        saved_template = cv2.imdecode(
            np.frombuffer(template_path.read_bytes(), np.uint8),
            cv2.IMREAD_COLOR,
        )
        pdf_path = _SCAN_DIR / "S25C-0i26090316340.pdf"
        current_templates = generate_ui_templates(
            str(pdf_path),
            config.page_count,
            config.rot_code,
            config.fine_angle,
            page_fine_angles=config.page_fine_angles,
        )
        remap = remap_preset_to_detected_layout(
            config,
            current_templates,
            source_templates={0: saved_template},
        )
        self.assertTrue(remap.accepted)

        from src.ui import MainWindow

        clean_pages = MainWindow._build_single_sample_template_pages(
            [saved_template],
            current_templates,
            remap.page_transforms,
            remap.config.page_count,
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            self.assertTrue(
                run_analysis(
                    [str(pdf_path)],
                    [current_templates[0]],
                    remap.config,
                    template_pages_preprocessed=True,
                    output_base_dir=temp_dir,
                    single_sample_template_pages=clean_pages,
                )
            )
            workbook_path = next(Path(temp_dir).rglob("*.xlsx"))
            workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            headers = [cell.value for cell in workbook["결과"][1]]
            values = [cell.value for cell in workbook["결과"][2]]

        answer = dict(zip(headers, values))
        self.assertEqual(answer["성별"], "남")
        self.assertEqual(answer["연령"], "50대")
        self.assertEqual(answer["경력"], "1년차")
        self.assertEqual([answer[f"Q{i}"] for i in range(1, 6)], [5] * 5)
        self.assertEqual(answer["의견"], "있음")


@unittest.skipUnless(
    _RUN_TRIO_CORPUS_TESTS,
    "set RUN_TRIO_CORPUS_TESTS=1 to run the independent-layout trio",
)
class IndependentFileLayoutCorpusTests(unittest.TestCase):
    """Exercise the real UI preset path and per-file analysis references."""

    @classmethod
    def setUpClass(cls):
        missing = [
            name for name in _INDEPENDENT_LAYOUT_TRIO
            if not (_SCAN_DIR / name).is_file()
        ]
        if missing:
            raise AssertionError(f"independent-layout corpus is missing: {missing}")
        preset_path = _PRESET_DIR / "기본.json"
        if not preset_path.is_file():
            raise AssertionError(f"default preset fixture is missing: {preset_path}")

    def test_trio_loads_default_preset_and_keeps_local_review_frames(self):
        """83 source pages must retain their own field-frame coordinates."""
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PyQt6.QtWidgets import QApplication

        from src.ui import MainWindow

        paths = [str(_SCAN_DIR / name) for name in _INDEPENDENT_LAYOUT_TRIO]
        app = QApplication.instance() or QApplication([])
        window = MainWindow()
        try:
            window.preset_dir = _PRESET_DIR
            with (
                patch(
                    "src.ui.QFileDialog.getOpenFileNames",
                    return_value=(paths, ""),
                ),
                patch("src.ui.QInputDialog.getInt", return_value=(1, True)),
                patch("src.ui.QMessageBox.information"),
                patch("src.ui.QMessageBox.critical") as error_dialog,
            ):
                self.assertTrue(window.load_pdf())
                self.assertFalse(error_dialog.called)

            window._load_preset_by_name("기본")
            self.assertEqual(window.current_preset_name, "기본")
            field_contract = [
                (
                    field.name,
                    tuple(field.value_map),
                    field.is_comment,
                    len(field.boxes),
                )
                for field in window.preset.fields
            ]
            self.assertEqual(
                [item[0] for item in field_contract],
                ["성별", "연령", "경력", "Q1", "Q2", "Q3", "Q4", "Q5", "의견"],
            )

            pages, preprocessed = window._analysis_input_pages()
            with tempfile.TemporaryDirectory() as temp_dir:
                self.assertTrue(
                    run_analysis(
                        paths,
                        pages,
                        window.preset,
                        template_pages_preprocessed=preprocessed,
                        single_sample_template_pages=(
                            window._single_sample_template_pages
                        ),
                        output_base_dir=temp_dir,
                    )
                )
                review_folder = next(Path(temp_dir).rglob("검토용"))
                expected_pages = dict(
                    zip(_INDEPENDENT_LAYOUT_TRIO, (19, 34, 30))
                )
                actual_total = 0
                expected_frames = sum(
                    item[3] for item in field_contract if not item[2]
                )
                self.assertEqual(expected_frames, 37)
                for name, page_count in expected_pages.items():
                    label = Path(name).stem
                    review_pdf = review_folder / f"{label}_원본포함.pdf"
                    self.assertTrue(review_pdf.is_file(), review_pdf)
                    document = fitz.open(review_pdf)
                    try:
                        self.assertEqual(len(document), page_count)
                        actual_total += len(document)
                        for page_idx in {0, min(1, page_count - 1), page_count - 1}:
                            self.assertGreaterEqual(
                                len(document[page_idx].get_drawings()),
                                expected_frames,
                            )
                            _assert_review_grid_alignment(
                                self, document, page_idx
                            )
                    finally:
                        document.close()
                self.assertEqual(actual_total, 83)

                # The source has two deliberately blank final 리더 sheets.
                # Review output keeps every scanned page, while Excel omits
                # only all-blank response rows.  Assert the exact remaining
                # labels so a filled respondent cannot disappear unnoticed.
                workbook_path = next(Path(temp_dir).rglob("*.xlsx"))
                workbook = openpyxl.load_workbook(
                    workbook_path, read_only=True, data_only=True
                )
                try:
                    sheet = workbook["결과"]
                    actual_rows = [
                        (row[0], row[1])
                        for row in sheet.iter_rows(min_row=2, values_only=True)
                    ]
                finally:
                    workbook.close()
                expected_rows = [
                    ("리더", f"리더_{page_idx}p")
                    for page_idx in range(1, 18)
                ]
                expected_rows.extend(
                    ("마음", f"마음_{page_idx}p")
                    for page_idx in range(1, 35)
                )
                expected_rows.extend(
                    ("여행", f"여행_{page_idx}p")
                    for page_idx in range(1, 31)
                )
                self.assertEqual(actual_rows, expected_rows)

            self.assertEqual(
                [
                    (
                        field.name,
                        tuple(field.value_map),
                        field.is_comment,
                        len(field.boxes),
                    )
                    for field in window.preset.fields
                ],
                field_contract,
            )
        finally:
            window._preset_dirty = False
            window.close()
            app.processEvents()


if __name__ == "__main__":
    unittest.main()
